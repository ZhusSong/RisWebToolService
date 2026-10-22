"""ASCII-only server tests. Run with the converter's Python dependencies."""
import base64
import io
import json
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).parents[1] / 'files'))
import converter_worker as w
import converter_server as s

def request(raw, suffix='pdf', fmt='docx'):
    return {'format': fmt, 'files': [{'name': 'input.' + suffix, 'data': base64.b64encode(raw).decode()}]}

def sample_pdf(scanned=False):
    from reportlab.pdfgen import canvas
    data = io.BytesIO()
    c = canvas.Canvas(data, pagesize=(500, 600))
    if not scanned:
        c.drawString(40, 550, 'Document conversion test')
        for y in (420, 450, 480):
            c.line(40, y, 340, y)
        for x in (40, 190, 340):
            c.line(x, 420, x, 480)
        for x, y, text in [(50, 460, 'Item'), (200, 460, 'Value'), (50, 430, '=1+1'), (200, 430, '42')]:
            c.drawString(x, y, text)
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    c.drawImage(ImageReader(Image.new('RGB', (80, 40), 'blue')), 40, 320, 80, 40)
    c.showPage(); c.save()
    return data.getvalue()

class ConverterTests(unittest.TestCase):
    def test_pdf_docx_keeps_text_table_image(self):
        from docx import Document
        result = w.convert(request(sample_pdf()))
        doc = Document(io.BytesIO(base64.b64decode(result['artifact']['data'])))
        self.assertIn('Document conversion test', '\n'.join(p.text for p in doc.paragraphs))
        self.assertTrue(doc.tables)
        self.assertTrue(doc.inline_shapes)

    def test_pdf_excel_cells_are_text_not_formulas(self):
        from openpyxl import load_workbook
        result = w.convert(request(sample_pdf(), fmt='xlsx'))
        book = load_workbook(io.BytesIO(base64.b64decode(result['artifact']['data'])))
        self.assertIn('Page1_Table1', book.sheetnames)
        self.assertEqual(book['Page1_Table1']['A2'].value, '=1+1')
        self.assertEqual(book['Page1_Table1']['A2'].data_type, 's')
        self.assertIn('Page1_Text', book.sheetnames)
        self.assertTrue(book['Page1_Preview']._images)

    def test_pdf_slides_and_scan_rejection(self):
        from pptx import Presentation
        raw = sample_pdf(scanned=True)
        for fmt in ('docx', 'xlsx'):
            with self.assertRaisesRegex(ValueError, 'scannedPdf'):
                w.convert(request(raw, fmt=fmt))
        result = w.convert(request(raw, fmt='pptx'))
        deck = Presentation(io.BytesIO(base64.b64decode(result['artifact']['data'])))
        self.assertEqual(len(deck.slides), 1)
        self.assertEqual(len(deck.slides[0].shapes), 1)

    def test_reject_page_limit_macros_and_external_resources(self):
        from pypdf import PdfWriter
        pdf = PdfWriter()
        for _ in range(21):
            pdf.add_blank_page(100, 100)
        data = io.BytesIO(); pdf.write(data)
        with self.assertRaisesRegex(ValueError, 'pageLimit'):
            w.convert(request(data.getvalue()))
        for name, content in [('word/vbaProject.bin', b'x'), ('word/_rels/document.xml.rels', b'<Relationships><Relationship TargetMode="External" Type="image" Target="http://example.com/a.png"/></Relationships>')]:
            data = io.BytesIO()
            with zipfile.ZipFile(data, 'w') as archive:
                archive.writestr(name, content)
            with self.assertRaises(ValueError):
                w.validate_office(data.getvalue(), '.docx')

    @unittest.skipUnless(shutil.which('libreoffice') or os.environ.get('SOFFICE_PATH'), 'LibreOffice is not installed')
    def test_office_pdf_real_engine(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        from pypdf import PdfReader
        doc = Document(); doc.add_heading('Word sample', 0)
        table = doc.add_table(rows=2, cols=2); table.cell(0,0).text = 'Table sample'
        book = Workbook(); book.active['A1'] = 'Excel sample'; book.active['B1'] = 42
        book.active.column_dimensions['A'].width = 24
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Slide sample'
        for suffix, document, expected in [('docx', doc, 'Word sample'), ('xlsx', book, 'Excel sample'), ('pptx', deck, 'Slide sample')]:
            raw = io.BytesIO(); document.save(raw)
            result = w.convert(request(raw.getvalue(), suffix, 'pdf'))
            pdf = PdfReader(io.BytesIO(base64.b64decode(result['artifact']['data'])))
            extracted = ''.join(page.extract_text() or '' for page in pdf.pages)
            # PDF glyph positioning can insert spaces inside words on Linux.
            # Ignore whitespace only; missing or reordered characters still fail.
            self.assertIn(''.join(expected.split()), ''.join(extracted.split()),
                          f'{suffix} PDF text content mismatch: {extracted!r}')

    @unittest.skipUnless(sys.platform == 'linux', 'Linux sandbox check')
    def test_sandbox_blocks_network(self):
        code = 'from converter_worker import deny_network; import socket; deny_network(); socket.socket(socket.AF_UNIX, socket.SOCK_STREAM).close(); socket.socket(socket.AF_INET, socket.SOCK_STREAM)'
        result = subprocess.run([sys.executable, '-c', code], cwd=Path(w.__file__).parent, capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(b'PermissionError', result.stderr)

    @unittest.skipUnless(sys.platform == 'linux' and shutil.which('libreoffice'), 'Linux conversion integration')
    def test_sandboxed_office_pipeline(self):
        from docx import Document
        from pypdf import PdfReader
        document = Document(); document.add_paragraph('Sandbox conversion')
        raw = io.BytesIO(); document.save(raw)
        result = s.run_job(request(raw.getvalue(), 'docx', 'pdf'))
        self.assertNotIn('error', result)
        pdf = PdfReader(io.BytesIO(base64.b64decode(result['artifact']['data'])))
        self.assertIn('Sandbox conversion', pdf.pages[0].extract_text())

    def test_server_origin_and_fixed_payload(self):
        server = s.ThreadingHTTPServer(('127.0.0.1', 0), s.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True); thread.start()
        try:
            with patch.object(s, 'run_job', return_value={'artifact': {'data':'x'}}) as run:
                req = urllib.request.Request(f'http://127.0.0.1:{server.server_port}/convert-api/process',
                    data=json.dumps({'format':'pdf','files':[], 'mode':'image','key':'unused'}).encode(), headers={'Origin':s.ORIGIN})
                with urllib.request.urlopen(req) as response:
                    self.assertEqual(response.status, 200)
                self.assertEqual(set(run.call_args.args[0]), {'format', 'files'})
                req.headers['Origin'] = 'https://example.org'
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(req)
                self.assertEqual(error.exception.code, 403); error.exception.close()
        finally:
            server.shutdown(); server.server_close(); thread.join()

if __name__ == '__main__':
    unittest.main()

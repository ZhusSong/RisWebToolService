"""Bounded Office/PDF conversion. Linux worker denies network socket creation."""
import base64
import contextlib
import ctypes
import errno
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from worker import MAX_FILE, OUTPUTS, DOCUMENTS, unpack, extract, artifact

OFFICE = {'.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx', '.odt', '.ods', '.odp'}
MAX_OUTPUT = 16 * 1024 * 1024
MAX_PAGES = 20

def deny_network():
    # Inherited by LibreOffice. Permit AF_UNIX for its private local IPC only.
    lib = ctypes.CDLL('libseccomp.so.2')
    class Compare(ctypes.Structure):
        _fields_ = [('arg', ctypes.c_uint), ('op', ctypes.c_int), ('a', ctypes.c_uint64), ('b', ctypes.c_uint64)]
    lib.seccomp_init.argtypes = [ctypes.c_uint32]
    lib.seccomp_init.restype = ctypes.c_void_p
    lib.seccomp_syscall_resolve_name.argtypes = [ctypes.c_char_p]
    lib.seccomp_rule_add_array.argtypes = [ctypes.c_void_p, ctypes.c_uint32, ctypes.c_int, ctypes.c_uint, ctypes.POINTER(Compare)]
    lib.seccomp_load.argtypes = [ctypes.c_void_p]
    lib.seccomp_release.argtypes = [ctypes.c_void_p]
    ctx = lib.seccomp_init(0x7fff0000)
    if not ctx:
        raise RuntimeError('sandboxUnavailable')
    try:
        # SCMP_CMP_NE: socket domain != AF_UNIX => EPERM.
        rule = Compare(0, 1, 1, 0)
        syscall = lib.seccomp_syscall_resolve_name(b'socket')
        if syscall < 0 or lib.seccomp_rule_add_array(ctx, 0x50000 | errno.EPERM, syscall, 1, ctypes.byref(rule)) != 0 or lib.seccomp_load(ctx) != 0:
            raise RuntimeError('sandboxUnavailable')
    finally:
        lib.seccomp_release(ctx)

def validate_office(raw, suffix):
    if suffix in {'.doc', '.xls', '.ppt'}:
        import olefile
        if not olefile.isOleFile(io.BytesIO(raw)):
            raise ValueError('invalid')
        with olefile.OleFileIO(io.BytesIO(raw)) as document:
            names = ['/'.join(parts).lower() for parts in document.listdir()]
            if any(any(word in name for word in ('vba', 'macros', 'objectpool', 'encryptedpackage')) for name in names):
                raise ValueError('unsafeOffice')
        return
    with unpack(raw, office=True) as archive:
        from defusedxml.ElementTree import fromstring
        for name in archive.namelist():
            if name.lower().startswith(('basic/', 'scripts/')) or '/externallinks/' in name.lower():
                raise ValueError('unsafeOffice')
            if name.endswith('.rels'):
                for link in fromstring(archive.read(name)):
                    if link.get('TargetMode') == 'External' and not link.get('Type', '').endswith('/hyperlink'):
                        raise ValueError('unsafeOffice')

def office_pdf(source, directory):
    profile = directory / 'profile'
    (profile / 'user').mkdir(parents=True)
    (profile / 'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
</oor:items>''', encoding='utf-8')
    executable = os.environ.get('SOFFICE_PATH', 'libreoffice')
    subprocess.run([executable, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo',
                    '--nodefault', '--norestore', '--convert-to', 'pdf', '--outdir', str(directory), str(source)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=85)
    output = source.with_suffix('.pdf')
    if not output.is_file() or output.stat().st_size > MAX_OUTPUT:
        raise ValueError('conversionFailed')
    return output.read_bytes()

def check_pdf(source):
    from pypdf import PdfReader
    pdf = PdfReader(source)
    if pdf.is_encrypted:
        raise ValueError('encrypted')
    if not 0 < len(pdf.pages) <= MAX_PAGES:
        raise ValueError('pageLimit')
    # Do not silently omit scanned pages in editable output.
    return all(bool((page.extract_text() or '').strip()) for page in pdf.pages)

def pdf_xlsx(source, target):
    import pdfplumber
    import pymupdf
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    cells = 0
    table_count = 0
    with pdfplumber.open(source) as pdf, pymupdf.open(source) as visual:
        for number, page in enumerate(pdf.pages, 1):
            preview = book.create_sheet(f'Page{number}_Preview')
            rendered = visual[number - 1]
            scale = min(1.5, 1200 / max(rendered.rect.width, rendered.rect.height))
            png = rendered.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).tobytes('png')
            preview.add_image(Image(io.BytesIO(png)), 'A1')
            preview.sheet_view.showGridLines = False
            tables = page.extract_tables()
            for index, rows in enumerate(tables, 1):
                table_count += 1
                sheet = book.create_sheet(f'Page{number}_Table{index}')
                sheet.freeze_panes = 'A2'
                for r, values in enumerate(rows, 1):
                    for c, value in enumerate(values, 1):
                        cells += 1
                        if cells > 20000:
                            raise ValueError('tableLimit')
                        cell = sheet.cell(r, c, value or '')
                        cell.data_type = 's'  # Never execute PDF text as a spreadsheet formula.
                        cell.alignment = Alignment(vertical='top', wrap_text=True)
                        if r == 1:
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill('solid', fgColor='E7EDF4')
                        sheet.column_dimensions[get_column_letter(c)].width = 24
            # Keep non-table text visible; it is not fabricated into table cells.
            text_sheet = book.create_sheet(f'Page{number}_Text')
            for row, line in enumerate((page.extract_text() or '').splitlines(), 1):
                if row > 10000:
                    raise ValueError('tableLimit')
                cell = text_sheet.cell(row, 1, line)
                cell.data_type = 's'
            text_sheet.column_dimensions['A'].width = 100
    book.save(target)
    return table_count

def pdf_pptx(source, target):
    import pymupdf
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    with pymupdf.open(source) as pdf:
        first = pdf[0].rect
        deck.slide_width = Inches(10)
        deck.slide_height = Inches(min(56, max(1, 10 * first.height / first.width)))
        for page in pdf:
            scale = min(1.5, 1600 / max(page.rect.width, page.rect.height))
            image = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).tobytes('png')
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            ratio = min(deck.slide_width / page.rect.width, deck.slide_height / page.rect.height)
            width, height = int(page.rect.width * ratio), int(page.rect.height * ratio)
            slide.shapes.add_picture(io.BytesIO(image), (deck.slide_width-width)//2, (deck.slide_height-height)//2, width=width, height=height)
    deck.save(target)

def convert(data):
    fmt, files = data.get('format'), data.get('files')
    if fmt not in OUTPUTS - {'text'} or not isinstance(files, list) or len(files) != 1 or not isinstance(files[0], dict):
        raise ValueError('invalid')
    item = files[0]
    if not isinstance(item.get('name'), str) or len(item['name']) > 200 or not isinstance(item.get('data'), str):
        raise ValueError('invalid')
    suffix = Path(item['name']).suffix.lower()
    raw = base64.b64decode(item['data'], validate=True)
    if not 0 < len(raw) <= MAX_FILE or suffix not in DOCUMENTS | OFFICE:
        raise ValueError('invalid')
    warning = ''
    with tempfile.TemporaryDirectory() as temporary:
        directory = Path(temporary)
        source = directory / ('source' + suffix)
        source.write_bytes(raw)
        target = directory / ('result.' + fmt)
        text_pages = check_pdf(source) if suffix == '.pdf' else False
        if suffix in OFFICE and fmt == 'pdf':
            validate_office(raw, suffix)
            result = office_pdf(source, directory)
            target.write_bytes(result)
            check_pdf(target)
        elif suffix == '.pdf' and fmt in {'docx', 'xlsx', 'pptx'}:
            if not text_pages and fmt != 'pptx':
                raise ValueError('scannedPdf')
            if fmt == 'docx':
                from pdf2docx import Converter
                converter = Converter(str(source))
                try:
                    converter.convert(str(target), multi_processing=False, ignore_page_error=False)
                finally:
                    converter.close()
            elif fmt == 'xlsx':
                if not pdf_xlsx(source, target):
                    warning = 'noTables'
            else:
                pdf_pptx(source, target)
        else:
            if suffix not in DOCUMENTS:
                raise ValueError('unsupportedPair')
            text = extract(item['name'], raw)
            if not text.strip() or len(text) > 60000:
                raise ValueError('invalid')
            return {'artifact': artifact(text, fmt)}
        if not target.is_file() or not 0 < target.stat().st_size <= MAX_OUTPUT:
            raise ValueError('outputLimit')
        return {'artifact': {'name': 'RISTOOL-result.' + fmt, 'data': base64.b64encode(target.read_bytes()).decode()}, 'warning': warning}

if __name__ == '__main__':
    try:
        if sys.platform != 'linux':
            raise RuntimeError('sandboxUnavailable')
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (95, 95))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT * 2, MAX_OUTPUT * 2))
        deny_network()
        with contextlib.redirect_stdout(sys.stderr):
            result = convert(json.load(sys.stdin))
    except ValueError as error:
        result = {'error': str(error) if str(error) in {'invalid', 'unsafeOffice', 'encrypted', 'pageLimit', 'scannedPdf', 'tableLimit', 'outputLimit', 'unsupportedPair'} else 'invalid'}
    except Exception:
        result = {'error': 'conversionFailed'}
    print(json.dumps(result))
